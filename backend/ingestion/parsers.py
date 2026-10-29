import csv
from io import BytesIO, StringIO
from pathlib import Path
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}


def parse_document(filename: str, content: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'no extension'}.")
    if suffix == ".pdf":
        pages = [{"text": page.extract_text() or "", "page": index + 1, "section": None} for index, page in enumerate(PdfReader(BytesIO(content)).pages)]
    elif suffix == ".docx":
        doc = Document(BytesIO(content))
        values = [p.text for p in doc.paragraphs if p.text.strip()]
        values += [" | ".join(cell.text.strip() for cell in row.cells) for table in doc.tables for row in table.rows]
        pages = [{"text": "\n".join(values), "page": None, "section": None}]
    elif suffix == ".csv":
        rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        pages = [{"text": "\n".join(" | ".join(row) for row in rows), "page": None, "section": "CSV"}]
    else:
        pages = [{"text": content.decode("utf-8-sig"), "page": None, "section": None}]
    cleaned = [{**item, "text": " ".join(item["text"].split())} for item in pages]
    result = [item for item in cleaned if item["text"]]
    if not result:
        raise ValueError("No text could be extracted. Scanned PDFs require OCR before upload.")
    return result
